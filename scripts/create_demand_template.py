from pathlib import Path

from openpyxl import Workbook


def main() -> None:
    output = Path("需求單資料夾") / "EMT技術模擬遊戲_需求單.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    data = workbook.active
    data.title = "_data"
    data.append(["資料類型", "功能編號", "功能名稱", "驗收條件", "優先級"])
    data.append(
        ["data_type", "feature_id", "feature_name", "acceptance_criteria", "priority"]
    )
    for row in [
        ["feature", "F001", "班次遊戲循環", "可完成三案並查看班次總評", "P0"],
        ["feature", "F002", "案例狀態與評分引擎", "行動會改變時間、穩定度與分數", "P0"],
        ["feature", "F003", "響應式介面", "桌機與手機均可完成遊戲", "P0"],
        ["content", "C001", "六個 EMT-1 案例", "六案皆可遊玩並附案後解析", "P0"],
        ["feature", "F004", "本機進度", "設定與最佳成績可於重新整理後保留", "P1"],
        ["deployment", "D001", "GitHub Pages", "正式建置可在子路徑載入", "P0"],
    ]:
        data.append(row)

    mapping = workbook.create_sheet("_mapping")
    mapping.append(["demand_script", "target", "note"])
    for row in [
        ["data_type", "implementation_route", "功能、內容或部署分類"],
        ["feature_id", "work_item_id", "唯一需求編號"],
        ["feature_name", "work_item_name", "顯示名稱"],
        ["acceptance_criteria", "test_requirement", "驗收條件"],
        ["priority", "delivery_priority", "MVP 優先級"],
    ]:
        mapping.append(row)

    metadata = workbook.create_sheet("_meta")
    metadata.append(["key", "value"])
    for row in [
        ["rule_file", "EMT技術模擬遊戲_規則.md"],
        ["pipeline_version", "2.0.0"],
        ["project_type", "SOFTWARE_PROJECT"],
        ["form_references", "[]"],
    ]:
        metadata.append(row)

    workbook.save(output)
    print(output)


if __name__ == "__main__":
    main()
